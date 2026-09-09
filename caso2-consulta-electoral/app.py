from datetime import datetime
from pathlib import Path
import os
import re
import socket

from flask import Flask, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


app = Flask(__name__)

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "datos"
DATA_DIR.mkdir(parents=True, exist_ok=True)

EXCEL_PATH = DATA_DIR / "miembros_de_mesa.xlsx"


def crear_excel():
    """Crea el archivo Excel si todavía no existe."""
    if EXCEL_PATH.exists():
        return

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Miembros de mesa"

    encabezados = [
        "DNI",
        "Miembro de mesa",
        "Nombres y apellidos",
        "Región",
        "Provincia",
        "Distrito",
        "Dirección del local de votación",
        "Fecha de registro",
    ]

    hoja.append(encabezados)

    fondo_encabezado = PatternFill(
        fill_type="solid",
        fgColor="028090",
    )

    for celda in hoja[1]:
        celda.font = Font(
            color="FFFFFF",
            bold=True,
        )
        celda.fill = fondo_encabezado
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    anchos = [14, 20, 30, 20, 20, 20, 45, 22]

    for indice, ancho in enumerate(anchos, start=1):
        columna = get_column_letter(indice)
        hoja.column_dimensions[columna].width = ancho

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = "A1:H1"

    libro.save(EXCEL_PATH)
    libro.close()


def obtener_registros():
    """Lee los registros almacenados en el Excel."""
    crear_excel()

    libro = load_workbook(EXCEL_PATH)
    hoja = libro["Miembros de mesa"]

    registros = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila[0]:
            registros.append(fila)

    libro.close()
    return registros


def dni_ya_registrado(dni):
    """Evita guardar dos veces el mismo DNI."""
    crear_excel()

    libro = load_workbook(EXCEL_PATH)
    hoja = libro["Miembros de mesa"]

    encontrado = any(
        str(fila[0].value).strip() == dni
        for fila in hoja.iter_rows(min_row=2)
        if fila[0].value is not None
    )

    libro.close()
    return encontrado


def guardar_registro(datos):
    """Guarda el resultado de la consulta en el Excel."""
    crear_excel()

    libro = load_workbook(EXCEL_PATH)
    hoja = libro["Miembros de mesa"]

    hoja.append(
        [
            datos["dni"],
            datos["miembro"],
            datos["nombres"],
            datos["region"],
            datos["provincia"],
            datos["distrito"],
            datos["direccion"],
            datetime.now().strftime("%d/%m/%Y %H:%M"),
        ]
    )

    ultima_fila = hoja.max_row

    for celda in hoja[ultima_fila]:
        celda.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    # El DNI se guarda como texto para conservar sus ocho dígitos.
    hoja.cell(row=ultima_fila, column=1).number_format = "@"

    libro.save(EXCEL_PATH)
    libro.close()


@app.route("/", methods=["GET", "POST"])
def inicio():
    mensaje = None
    error = None

    if request.method == "POST":
        dni = request.form.get("dni", "").strip()
        miembro = request.form.get("miembro", "").strip()
        nombres = request.form.get("nombres", "").strip()
        region = request.form.get("region", "").strip()
        provincia = request.form.get("provincia", "").strip()
        distrito = request.form.get("distrito", "").strip()
        direccion = request.form.get("direccion", "").strip()

        if not re.fullmatch(r"\d{8}", dni):
            error = "El DNI debe contener exactamente 8 números."

        elif miembro not in {"Sí", "No"}:
            error = "Selecciona si la persona es miembro de mesa."

        elif dni_ya_registrado(dni):
            error = "Este DNI ya fue registrado."

        elif miembro == "Sí" and not all(
            [nombres, region, provincia, distrito, direccion]
        ):
            error = (
                "Si es miembro de mesa, completa los nombres, "
                "la ubicación y la dirección del local."
            )

        else:
            if miembro == "No":
                nombres = "No corresponde"
                region = "No corresponde"
                provincia = "No corresponde"
                distrito = "No corresponde"
                direccion = "No corresponde"

            guardar_registro(
                {
                    "dni": dni,
                    "miembro": miembro,
                    "nombres": nombres,
                    "region": region,
                    "provincia": provincia,
                    "distrito": distrito,
                    "direccion": direccion,
                }
            )

            mensaje = "Registro guardado correctamente en el archivo Excel."

    registros = obtener_registros()

    return render_template(
        "index.html",
        mensaje=mensaje,
        error=error,
        registros=registros,
        hostname=socket.gethostname(),
    )


@app.route("/descargar-excel")
def descargar_excel():
    """Permite descargar el archivo generado."""
    crear_excel()

    return send_file(
        EXCEL_PATH,
        as_attachment=True,
        download_name="miembros_de_mesa.xlsx",
    )


@app.route("/health")
def health():
    """Permite que Docker compruebe el estado de la aplicación."""
    return {"estado": "saludable"}, 200


if __name__ == "__main__":
    puerto = int(os.environ.get("PORT", 5000))

    app.run(
        host="0.0.0.0",
        port=puerto,
    )